import os
import re
import sys

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import documents  # noqa: E402


MODEL = "gpt-4"


def test_chunk_budget_includes_paragraph_separators():
    chunks = documents.chunk_text("hello\n\nworld", 2, MODEL)
    assert all(documents.count_tokens(chunk, MODEL) <= 2 for chunk in chunks)
    assert " ".join(chunks) == "hello world"


@pytest.mark.parametrize("text", ["bằng tiếng Việt " * 30, "😀" * 30])
@pytest.mark.parametrize("budget", [3, 5, 7])
def test_chunking_preserves_unicode_characters(text, budget):
    chunks = documents.chunk_text(text, budget, MODEL)
    assert all(documents.count_tokens(chunk, MODEL) <= budget for chunk in chunks)
    assert "\ufffd" not in "".join(chunks)
    assert re.sub(r"\s", "", "".join(chunks)) == re.sub(r"\s", "", text)


def test_chunking_reports_budget_too_small_for_a_unicode_character():
    with pytest.raises(ValueError, match="complete Unicode character"):
        documents.chunk_text("😀😀", 1, MODEL)


def test_token_count_and_chunking_accept_literal_special_token_text():
    text = "A document describing <|endoftext|> and <|fim_prefix|> tokens."
    expected = len(documents.get_encoding(MODEL).encode(text, disallowed_special=()))
    assert documents.count_tokens(text, MODEL) == expected
    chunks = documents.chunk_text(text, 8, MODEL)
    assert all(documents.count_tokens(chunk, MODEL) <= 8 for chunk in chunks)
    assert re.sub(r"\s", "", "".join(chunks)) == re.sub(r"\s", "", text)


def test_word_extraction_keeps_tables_in_document_order(tmp_path):
    path = tmp_path / "lesson.docx"
    document = Document()
    document.add_paragraph("Bảng một")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First table"
    table.cell(0, 1).text = "Nội dung một"
    document.add_paragraph("Bảng hai")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Second table"
    document.add_paragraph("Kết thúc")
    document.save(path)

    text = documents.extract_text(path, use_cache=False)
    assert text.splitlines() == [
        "Bảng một", "First table | Nội dung một", "Bảng hai", "Second table", "Kết thúc",
    ]


@pytest.mark.parametrize("as_list", [False, True])
def test_document_context_accepts_path_objects(tmp_path, as_list):
    path = tmp_path / "lesson.txt"
    path.write_text("Nội dung tiếng Việt", encoding="utf-8")
    text, problems = documents.build_document_context([path] if as_list else path)
    assert problems == []
    assert "File: lesson.txt" in text
    assert "Nội dung tiếng Việt" in text


def test_empty_powerpoint_is_reported_instead_of_using_slide_labels_as_content(tmp_path):
    path = tmp_path / "empty.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(3), Inches(1))
    presentation.save(path)

    with pytest.raises(documents.DocumentError, match="No text found"):
        documents.extract_text(path, use_cache=False)


def test_powerpoint_reads_nested_groups_and_tolerates_missing_notes_placeholder(tmp_path):
    path = tmp_path / "grouped.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape().shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Nội dung nhóm"
    notes = slide.notes_slide.notes_placeholder
    notes.element.getparent().remove(notes.element)
    presentation.save(path)

    text = documents.extract_text(path, use_cache=False)
    assert "--- Slide 1 ---" in text
    assert "Nội dung nhóm" in text
