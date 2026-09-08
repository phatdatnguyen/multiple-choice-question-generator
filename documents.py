"""Document text extraction, token counting and chunking."""

import hashlib
import os
import re
import zipfile
from xml.etree import ElementTree

import html2text
import tiktoken

try:  # pypdf is the maintained successor to PyPDF2
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - fallback for older installs
    from PyPDF2 import PdfReader

from docx import Document
from pptx import Presentation

SUPPORTED_EXTENSIONS = [
    ".pdf", ".docx", ".pptx", ".htm", ".html", ".epub",
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".rst",
]

# Formats python-docx / python-pptx cannot open, reported with advice instead of
# a confusing parser traceback.
LEGACY_EXTENSIONS = {
    ".doc": "Word 97-2003",
    ".ppt": "PowerPoint 97-2003",
    ".xls": "Excel 97-2003",
    ".rtf": "Rich Text",
    ".pages": "Apple Pages",
    ".key": "Apple Keynote",
}

_TEXT_CACHE = {}
_CACHE_LIMIT = 32


class DocumentError(Exception):
    """Raised when a document cannot be read."""


# --------------------------------------------------------------------------- #
# Extractors
# --------------------------------------------------------------------------- #

def read_pdf_file(file_path):
    reader = PdfReader(file_path)
    # extract_text() returns None for pages with no text layer.
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def read_word_file(file_path):
    document = Document(file_path)
    # Keep tables next to the paragraphs that explain them. Reading all
    # paragraphs first would move every table to the end of the document.
    content = {paragraph._element: [paragraph.text] for paragraph in document.paragraphs}
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        content[table._element] = rows
    parts = [part for element in document.element.body for part in content.get(element, [])]
    return "\n".join(parts)


def read_powerpoint_file(file_path):
    presentation = Presentation(file_path)
    parts = []
    for number, slide in enumerate(presentation.slides, start=1):
        slide_parts = list(_powerpoint_shape_text(slide.shapes))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                slide_parts.append(f"[Notes] {notes.text}")
        if slide_parts:
            parts.append(f"--- Slide {number} ---")
            parts.extend(slide_parts)
    return "\n".join(parts)


def _powerpoint_shape_text(shapes):
    for shape in shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            yield shape.text_frame.text
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    yield " | ".join(cells)
        elif hasattr(shape, "shapes"):
            yield from _powerpoint_shape_text(shape.shapes)


def _html_to_text(markup):
    converter = html2text.HTML2Text()
    converter.ignore_links = True
    converter.ignore_images = True
    converter.body_width = 0
    return converter.handle(markup)


def read_html_file(file_path):
    return _html_to_text(_read_raw_text(file_path))


def read_epub_file(file_path):
    """Extract an EPUB by reading its spine in reading order."""
    with zipfile.ZipFile(file_path) as archive:
        names = archive.namelist()
        opf_name = next((name for name in names if name.lower().endswith(".opf")), None)

        hrefs = []
        if opf_name:
            root = ElementTree.fromstring(archive.read(opf_name))
            namespace = {"opf": "http://www.idpf.org/2007/opf"}
            manifest = {
                item.get("id"): item.get("href")
                for item in root.iterfind(".//opf:manifest/opf:item", namespace)
            }
            base = os.path.dirname(opf_name)
            for itemref in root.iterfind(".//opf:spine/opf:itemref", namespace):
                href = manifest.get(itemref.get("idref"))
                if href:
                    joined = os.path.normpath(os.path.join(base, href)).replace("\\", "/")
                    if joined in names:
                        hrefs.append(joined)

        if not hrefs:  # No usable spine - fall back to every document in the zip.
            hrefs = sorted(
                name for name in names
                if name.lower().endswith((".xhtml", ".html", ".htm"))
            )

        parts = []
        for href in hrefs:
            markup = archive.read(href).decode("utf-8", errors="replace")
            parts.append(_html_to_text(markup))
    return "\n".join(parts)


def _read_raw_text(file_path):
    """Read a text file, tolerating imperfect encodings rather than crashing."""
    with open(file_path, "rb") as file:
        data = file.read()

    # UTF-16 is only tried when a byte-order mark says so: it decodes almost any
    # even-length input into plausible-looking mojibake otherwise.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeError:
            pass

    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_text_file(file_path):
    return _read_raw_text(file_path)


_EXTRACTORS = {
    ".pdf": read_pdf_file,
    ".docx": read_word_file,
    ".pptx": read_powerpoint_file,
    ".htm": read_html_file,
    ".html": read_html_file,
    ".epub": read_epub_file,
}


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #

def _file_digest(file_path):
    digest = hashlib.sha256()
    with open(file_path, "rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_text(file_path, *, use_cache=True):
    """Extract plain text from one document.

    Results are cached by file content, so switching models and regenerating
    does not re-parse a large PDF.
    """
    extension = os.path.splitext(file_path)[1].lower()
    if extension in LEGACY_EXTENSIONS:
        raise DocumentError(
            f"{LEGACY_EXTENSIONS[extension]} files ({extension}) are not supported. "
            f"Please save the file as .docx, .pptx or .pdf and try again."
        )

    key = None
    if use_cache:
        try:
            key = f"{_file_digest(file_path)}:{extension}"
        except OSError:
            key = None
        if key and key in _TEXT_CACHE:
            return _TEXT_CACHE[key]

    extractor = _EXTRACTORS.get(extension, read_text_file)
    try:
        text = extractor(file_path)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(f"Could not read {os.path.basename(file_path)}: {exc}") from exc

    text = (text or "").strip()
    if not text:
        raise DocumentError(
            f"No text found in {os.path.basename(file_path)}. "
            f"If it is a scanned PDF it needs OCR first."
        )

    if key:
        if len(_TEXT_CACHE) >= _CACHE_LIMIT:
            _TEXT_CACHE.pop(next(iter(_TEXT_CACHE)))
        _TEXT_CACHE[key] = text
    return text


def build_document_context(file_paths, *, use_cache=True):
    """Extract several documents into one delimited string.

    Returns ``(text, problems)``; unreadable files are reported but do not stop
    the readable ones from being used.
    """
    if isinstance(file_paths, (str, bytes, os.PathLike)) or file_paths is None:
        file_paths = [file_paths] if file_paths else []

    sections, problems = [], []
    for file_path in file_paths:
        # pathlib.Path.name is only the basename, while uploaded file objects
        # expose their full temporary path through .name.
        path = (
            file_path if isinstance(file_path, (str, bytes, os.PathLike))
            else getattr(file_path, "name", file_path)
        )
        try:
            text = extract_text(path, use_cache=use_cache)
        except DocumentError as exc:
            problems.append(str(exc))
            continue
        sections.append(
            "<<<DOCUMENT_CONTENT>>>\n"
            f"File: {os.path.basename(path)}\n"
            f"{text}\n"
            "<<<END_DOCUMENT>>>"
        )
    return "\n\n".join(sections), problems


# --------------------------------------------------------------------------- #
# Tokens
# --------------------------------------------------------------------------- #

def get_encoding(model):
    try:
        return tiktoken.encoding_for_model(model)
    except (KeyError, ValueError):
        # Newer models share the o200k_base vocabulary.
        try:
            return tiktoken.get_encoding("o200k_base")
        except (KeyError, ValueError):
            return tiktoken.get_encoding("cl100k_base")


def count_tokens(text, model):
    """Count tokens in a string."""
    return len(get_encoding(model).encode(str(text or ""), disallowed_special=()))


def chunk_text(text, max_tokens, model):
    """Split text into chunks of at most ``max_tokens`` tokens.

    Splits on blank lines first, then single lines, and only falls back to
    slicing tokens for a single line that is itself too long.
    """
    if max_tokens <= 0:
        raise ValueError("max_tokens must be greater than 0.")

    encoding = get_encoding(model)
    text = str(text or "")
    if len(encoding.encode(text, disallowed_special=())) <= max_tokens:
        return [text] if text.strip() else []

    chunks, current, current_tokens = [], [], 0
    separator_tokens = len(encoding.encode("\n\n", disallowed_special=()))

    def flush():
        nonlocal current, current_tokens
        if current:
            # Re-tokenize the actual result: joining/stripping can change token
            # boundaries, so the sum of the individual counts is only an estimate.
            chunks.extend(_split_token_blocks("\n\n".join(current).strip(), max_tokens, encoding))
            current, current_tokens = [], 0

    for block in _split_blocks(text, max_tokens, encoding):
        block_tokens = len(encoding.encode(block, disallowed_special=()))
        if current and current_tokens + separator_tokens + block_tokens > max_tokens:
            flush()
        if current:
            current_tokens += separator_tokens
        current.append(block)
        current_tokens += block_tokens
    flush()
    return [chunk for chunk in chunks if chunk.strip()]


def _split_blocks(text, max_tokens, encoding):
    """Yield blocks that each fit within ``max_tokens``."""
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(encoding.encode(paragraph, disallowed_special=())) <= max_tokens:
            yield paragraph
            continue
        for line in paragraph.splitlines():
            line = line.strip()
            if not line:
                continue
            yield from _split_token_blocks(line, max_tokens, encoding)


def _split_token_blocks(text, max_tokens, encoding):
    """Split at token boundaries that also preserve complete UTF-8 characters."""
    tokens = encoding.encode(text, disallowed_special=())
    start = 0
    while start < len(tokens):
        end = min(start + max_tokens, len(tokens))
        while end > start:
            try:
                block = encoding.decode(tokens[start:end], errors="strict")
            except UnicodeDecodeError:
                end -= 1
                continue
            if len(encoding.encode(block, disallowed_special=())) <= max_tokens:
                break
            end -= 1
        if end == start:
            raise ValueError("max_tokens is too small to contain a complete Unicode character.")
        yield block
        start = end
